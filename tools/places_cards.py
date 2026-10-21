"""
База мест Voyz в Excel в формате «одна карточка — один лист».

  python tools/places_cards.py export [data/voyz-places.xlsx]   # places.json + переводы → Excel
  python tools/places_cards.py check  [data/voyz-places.xlsx]   # проверки, ничего не пишет
  python tools/places_cards.py import [data/voyz-places.xlsx]   # Excel → places.json + translations/en
  python tools/places_cards.py roundtrip                        # export → import → сравнение с places.json

Лист карточки называется её id (kg-skazka). Внутри — разделы в порядке карточки приложения.
Колонка A — машинная разметка для парсера, B — подпись, C — значение (для таблиц C, D, E… — колонки).
  • строка с A = ключ поля (name, short, …)       → значение поля в C;
  • A = «#section»                                → заголовок раздела (для людей);
  • A = «#table:<имя>»                            → начало таблицы-списка; следом «#keys» (ключи колонок с C) и «#labels» (подписи);
    строки данных — с пустой A, пустые строки пропускаются; таблица заканчивается на следующей строке с непустой A.
Листы с другими именами (README, Оглавление, теги, тикеты, справочники, _шаблон) парсер карточек игнорирует.
Схема полей и проверки — общие с tools/places_xlsx.py.
"""
import json
import os
import re
import sys
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import places_xlsx as X  # noqa: E402

ID_RX = re.compile(r'(kg|alt)-[a-z0-9-]+')
FIELD = {c[0]: c for c in X.PLACE_COLS}  # ключ → (ключ, путь, подпись, ширина, обязательное)
CHILD = {c[0]: c for c in X.CHILD}       # имя → (имя, подпись, путь, поля)
ROUTE_EXTRA = [('has_rt', 'Есть маршрут на карте (да/нет)', 10)]
TR_TABLE = ('translations_en', 'Перевод на английский', None, [('path', 'Поле', 22), ('ru', 'Русский (только чтение)', 60), ('en', 'English', 60)])

# разделы в порядке карточки приложения: (название, поля, таблицы)
SECTIONS = [
    ('Основное', ['id', 'country', 'publish', 'status', 'name', 'full', 'cat', 'reg', 'dist', 'lat', 'lon', 'elev', 'near', 'tags', 'top', 'rank', 'updated_at', 'updated_by'], []),
    ('Коротко на карточке (плашки)', ['months', 's_diff', 's_when', 's_acc', 's_onsite', 's_entry', 's_net', 'rating_v', 'rating_n', 'rating_src'], []),
    ('Обзор — тексты', ['short', 'desc', 'guide', 'see', 'skip'], []),
    ('Медиа', [], ['photos', 'videos', 'photo_links']),
    ('Дорога', ['how', 'needcar', 'self', 'roadtype', 'traveltime', 'transfer', 'p_sum'], ['routes', 'transport']),
    ('Сезон и что взять', ['when', 'diff', 'hours', 'daytime', 'p_season', 'bring'], []),
    ('Важно', ['warn', 'entry', 'net', 'p_permits', 'p_access'], []),
    ('Ночлег и еда', ['stay', 'food'], ['stays', 'food']),
    ('На месте', ['camping', 'photo_spots'], ['hikes', 'horses', 'activities', 'rent', 'itinerary']),
    ('Сервисы рядом', [k for k in FIELD if k.startswith('svc_')], []),
    ('Бюджет', ['budget_economy', 'budget_comfort', 'budget_year'], []),
    ('Что говорят путешественники', ['reviews'], []),
    ('Рядом', [], ['nearby']),
    ('Проверка и качество данных', ['checked', 'note', 'gaps', 'flags', 'idx', 'rich', 'conf', 'prac', 'en_reviewed'], ['fixes', 'sources', 'sources_practical']),
    ('Перевод EN', [], ['translations_en']),
]
_covered = [f for _, fs, _ in SECTIONS for f in fs]
assert sorted(_covered) == sorted(FIELD), set(FIELD) ^ set(_covered)

# стиль
F = 'Arial'
C_INK, C_LIME, C_MIST = '111111', 'D8FC00', 'F1F1EF'
KEY = Font(name=F, size=8, color='9A9A9A')
LBL = Font(name=F, size=10, bold=True)
VAL = Font(name=F, size=10)
TITLE = Font(name=F, size=18, bold=True)
SEC_FONT = Font(name=F, size=12, bold=True, color=C_LIME)
SEC_FILL = PatternFill('solid', fgColor=C_INK)
TBL_FONT = Font(name=F, size=10, bold=True)
TBL_FILL = PatternFill('solid', fgColor='DEEAD2')
REQ_FILL = PatternFill('solid', fgColor='FFF2A8')
LBL_FILL = PatternFill('solid', fgColor=C_MIST)
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
LONG = {'short', 'desc', 'guide', 'see', 'skip', 'how', 'warn', 'p_sum', 'p_season', 'when', 'diff', 'stay', 'food', 'reviews', 'checked', 'note', 'gaps', 'flags', 'p_access', 'p_permits'}
YN_FIELDS = {'top', 'en_reviewed'}
LIST_FIELDS = {'country': 'country', 'publish': 'publish', 'status': 'status', 'cat': 'cat', 'reg': 'reg', 's_acc': 's_acc', 's_diff': 's_diff', 's_when': 's_when', 'top': 'yes_no', 'en_reviewed': 'yes_no'}
STRICT = {'country', 'publish', 'top', 'en_reviewed'}


def table_fields(name):
    if name == 'translations_en':
        return TR_TABLE[3]
    f = list(CHILD[name][3])
    return f + ROUTE_EXTRA if name == 'routes' else f


def table_title(name):
    return TR_TABLE[1] if name == 'translations_en' else CHILD[name][1]


# ---------------------------------------------------------------- запись одной карточки
class Card:
    def __init__(self, ws, refs):
        self.ws, self.r, self.refs = ws, 1, refs
        self.dv = {}

    def cell(self, c, v, font=VAL, fill=None, align=TOP):
        x = self.ws.cell(self.r, c, v)
        x.font = font
        x.alignment = align
        if fill:
            x.fill = fill
        return x

    def section(self, title):
        self.r += 1
        self.cell(1, '#section', KEY)
        for c in range(2, 9):
            self.ws.cell(self.r, c).fill = SEC_FILL
        self.cell(2, title, SEC_FONT, SEC_FILL)
        self.ws.row_dimensions[self.r].height = 22
        self.r += 1

    def field(self, key, value):
        _, _, label, _, req = FIELD[key]
        self.cell(1, key, KEY)
        self.cell(2, label, LBL, REQ_FILL if req else LBL_FILL, WRAP)
        c = self.cell(3, value, VAL, None, WRAP if key in LONG else TOP)
        if key in LIST_FIELDS:
            self.validate(f'C{self.r}', LIST_FIELDS[key], key in STRICT)
        self.r += 1
        return c

    def table(self, name, rows):
        fields = table_fields(name)
        self.cell(1, f'#table:{name}', KEY)
        self.cell(2, f'{table_title(name)} — {len(rows)}', TBL_FONT, TBL_FILL)
        self.r += 1
        self.cell(1, '#keys', KEY)
        for j, (k, *_) in enumerate(fields):
            self.cell(3 + j, k, KEY)
        self.r += 1
        self.cell(1, '#labels', KEY)
        self.cell(2, '№ / строки ниже', Font(name=F, size=8, italic=True, color='9A9A9A'))
        for j, (_, lbl, _) in enumerate(fields):
            self.cell(3 + j, lbl, LBL, LBL_FILL, WRAP)
        self.r += 1
        for row in rows:
            self.cell(2, row.get('pos'), Font(name=F, size=9, color='6E6E6A'))
            for j, (k, *_) in enumerate(fields):
                v = row.get(k)
                self.cell(3 + j, v, VAL, None, WRAP if isinstance(v, str) and len(v) > 60 else TOP)
            self.r += 1
        self.r += 1  # пустая строка — место для новой записи (её можно и вставить)

    def validate(self, ref, list_name, strict):
        key = (list_name, strict)
        if key not in self.dv:
            dv = DataValidation(type='list', formula1=self.refs[list_name], allow_blank=True, showErrorMessage=True,
                                errorStyle='stop' if strict else 'warning', errorTitle='Значение не из справочника', error='Выберите из списка (лист «Справочники»)')
            self.ws.add_data_validation(dv)
            self.dv[key] = dv
        self.dv[key].add(ref)


def write_card(wb, sheet_name, p, row, children, tr_rows, refs):
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.tabColor = '3F7D00' if row.get('country') == 'kg' else '1A4DFF'
    for col, w in zip('ABCDEFGHIJKLMNOPQRSTUVWXYZ', [11, 34, 60, 30, 24, 22, 22, 22] + [18] * 19):
        ws.column_dimensions[col].width = w
    card = Card(ws, refs)
    t = ws.cell(1, 2, row.get('name') or 'Новое место')
    t.font = TITLE
    ws.cell(1, 1, '#title').font = KEY
    ws.cell(1, 3, f"{row.get('id') or ''} · ← к оглавлению").font = Font(name=F, size=10, color='1A4DFF', underline='single')
    ws.cell(1, 3).hyperlink = "#'Оглавление'!A1"
    ws.row_dimensions[1].height = 28
    for title, fields, tables in SECTIONS:
        card.section(title)
        for k in fields:
            card.field(k, row.get(k))
        for tname in tables:
            card.table(tname, tr_rows if tname == 'translations_en' else children.get(tname, []))
    ws.freeze_panes = 'C2'
    return ws


# ---------------------------------------------------------------- экспорт
def export(out):
    data = json.load(open(X.PLACES, encoding='utf-8'))
    P = data['places']
    revf = os.path.join(X.TR_EN, '_reviewed.json')
    reviewed = set(json.load(open(revf, encoding='utf-8'))) if os.path.exists(revf) else set()
    old = read_extras(out)

    wb = Workbook()
    wb.remove(wb.active)
    readme(wb.create_sheet('README'), len(P))
    index_ws = wb.create_sheet('Оглавление')
    tags_ws = wb.create_sheet('Теги')
    tickets_ws = wb.create_sheet('Тикеты')
    lists_ws = wb.create_sheet('Справочники')
    refs = write_lists(lists_ws, P)

    # шаблон новой карточки
    tpl = write_card(wb, '_шаблон', None, {'publish': 'черновик', 'status': 'Нужна проверка'}, {}, [], refs)
    tpl.sheet_properties.tabColor = 'D8FC00'

    idx_rows = []
    for p in P:
        tags = old['tags'].get(p['id']) or p.get('tags') or X.auto_tags(p)
        row = X.place_row(p, tags, p['id'] in reviewed)
        children = {}
        for name, _, path, _ in X.CHILD:
            rows = X.child_rows(name, p, X.gp(p, path) or [])
            if name == 'routes':
                for r in rows:
                    r['has_rt'] = X.yn(r.pop('_rt'))
            children[name] = rows
        ru = json.load(open(os.path.join(X.TR_RU, f"{p['id']}.json"), encoding='utf-8'))
        enf = os.path.join(X.TR_EN, f"{p['id']}.json")
        en = json.load(open(enf, encoding='utf-8')) if os.path.exists(enf) else {}
        tr_rows = [dict(pos=i, path=k, ru=v, en=en.get(k)) for i, (k, v) in enumerate(ru.items(), 1)]
        write_card(wb, p['id'], p, row, children, tr_rows, refs)
        idx_rows.append((p, row, len(children['photos']), len(children['routes']), len(children['stays'])))

    write_index(index_ws, idx_rows)
    write_simple(tags_ws, ['tag_id', 'name_ru', 'name_en', 'group', 'description'], ['ID тега (латиница, дефис)', 'Название RU', 'Название EN', 'Группа', 'Описание'],
                 old['tag_dict'] or X.tag_dict(), [18, 26, 26, 14, 50], {'tag_id', 'name_ru'})
    write_simple(tickets_ws, ['ticket_id', 'type', 'place_id', 'title', 'details', 'status', 'priority', 'assignee', 'created', 'due', 'source', 'done_at'],
                 ['№ тикета', 'Тип', 'Карточка (id)', 'Задача', 'Подробности', 'Статус', 'Приоритет', 'Исполнитель', 'Создан', 'Срок', 'Откуда задача', 'Закрыт'],
                 old['tickets'] or X.seed_tickets(P, reviewed), [10, 18, 14, 40, 60, 12, 10, 14, 11, 11, 18, 11], {'ticket_id', 'type', 'title', 'status'})
    for col, lst, strict in (('B', 'ticket_type', True), ('F', 'ticket_status', True), ('G', 'priority', True)):
        dv = DataValidation(type='list', formula1=refs[lst], allow_blank=True, showErrorMessage=True, errorStyle='stop')
        tickets_ws.add_data_validation(dv)
        dv.add(f'{col}3:{col}3000')
    # ссылки из тикетов на карточки
    for r in range(3, tickets_ws.max_row + 1):
        c = tickets_ws.cell(r, 3)
        if c.value and c.value in wb.sheetnames:
            c.hyperlink = f"#'{c.value}'!A1"
            c.font = Font(name=F, size=10, color='1A4DFF', underline='single')

    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'✓ {os.path.relpath(out, X.ROOT)}: карточек {len(P)}, листов {len(wb.sheetnames)}')


def write_simple(ws, keys, labels, rows, widths, required):
    X.write_table(ws, keys, labels, rows, widths, required=required, freeze='B3')


def write_lists(ws, P):
    cols = {
        'country': ['kg', 'alt'], 'publish': X.PUBLISH, 'status': sorted({p['status'] for p in P} | {'Нужна проверка'}),
        'cat': sorted({p['cat'] for p in P}), 'reg': sorted({p['reg'] for p in P}), 's_acc': sorted({p['s']['acc'] for p in P}),
        's_diff': sorted({p['s']['diff'] for p in P}), 's_when': sorted({p['s']['when'] for p in P}), 'yes_no': [X.YES, X.NO],
        'ticket_type': X.TICKET_TYPES, 'ticket_status': X.TICKET_STATUS, 'priority': X.PRIORITY,
        'license': sorted({ph['lic'] for p in P for ph in p['photos']}),
    }
    for i, (k, vals) in enumerate(cols.items(), 1):
        ws.cell(1, i, k).font = KEY
        ws.cell(2, i, k).font = LBL
        ws.column_dimensions[get_column_letter(i)].width = 24
        for r, v in enumerate(vals, 3):
            ws.cell(r, i, v).font = VAL
    return {k: f"'Справочники'!${get_column_letter(i)}$3:${get_column_letter(i)}${len(v) + 2}" for i, (k, v) in enumerate(cols.items(), 1)}


def find_row(wb_or_ws, key):
    """Номер строки поля в карточке (одинаков у всех карточек — раздел «Основное» фиксирован)."""
    r = 1
    for title, fields, _ in SECTIONS[:1]:
        r += 2  # пустая + заголовок раздела
        for k in fields:
            if k == key:
                return r
            r += 1
    raise KeyError(key)


def write_index(ws, rows):
    keys = ['id', 'name', 'country', 'publish', 'status', 'cat', 'reg', 'photos', 'routes', 'stays']
    labels = ['Карточка (ссылка)', 'Название', 'Страна', 'Публикация', 'Статус', 'Категория', 'Регион', 'Фото', 'Подъездов', 'Ночлег']
    X.style_header(ws, keys, labels, (), [18, 34, 8, 14, 20, 24, 24, 7, 9, 8])
    for i, (p, row, nph, nrt, nst) in enumerate(rows, 3):
        sid = p['id']
        c = ws.cell(i, 1, sid)
        c.hyperlink = f"#'{sid}'!A1"
        c.font = Font(name=F, size=10, color='1A4DFF', underline='single')
        # значения подтягиваются формулой из карточки — правки в карточке сразу видны здесь
        for j, k in enumerate(['name', 'country', 'publish', 'status', 'cat', 'reg'], 2):
            x = ws.cell(i, j, f"='{sid}'!C{find_row(None, k)}")
            x.font = VAL
        for j, v in enumerate([nph, nrt, nst], 8):
            ws.cell(i, j, v).font = VAL
    n = len(rows) + 2
    ws.freeze_panes = 'B3'
    ws.auto_filter.ref = f'A2:J{n}'
    # сводка справа
    sx = 12
    ws.column_dimensions[get_column_letter(sx)].width = 30
    ws.column_dimensions[get_column_letter(sx + 1)].width = 10
    summary = [('Карточек', f'=COUNTA(A3:A{n})'), ('Кыргызстан', f'=COUNTIF(C3:C{n},"kg")'), ('Алтай', f'=COUNTIF(C3:C{n},"alt")'),
               ('Опубликовано', f'=COUNTIF(D3:D{n},"опубликовано")'), ('Черновиков', f'=COUNTIF(D3:D{n},"черновик")'),
               ('Без фото', f'=COUNTIF(H3:H{n},0)'), ('Открытых тикетов', "=COUNTIF('Тикеты'!F:F,\"открыт\")"),
               ('В работе', "=COUNTIF('Тикеты'!F:F,\"в работе\")")]
    ws.cell(2, sx, 'Сводка').font = Font(name=F, size=12, bold=True)
    for k, (a, f) in enumerate(summary, 3):
        ws.cell(k, sx, a).font = VAL
        ws.cell(k, sx + 1, f).font = Font(name=F, size=10, bold=True)
    ws.cell(len(summary) + 4, sx, 'Фото/подъезды/ночлег — на момент выгрузки').font = Font(name=F, size=8, italic=True, color='9A9A9A')


def readme(ws, count):
    lines = [
        ('База мест Voyz — одна карточка на лист', 'h'),
        (f'Выгружено {date.today().isoformat()} · карточек: {count}', ''),
        ('', ''),
        ('Как пользоваться', 'h2'),
        ('«Оглавление» — список всех карточек со ссылками (клик по id открывает лист). Сводка справа.', ''),
        ('Лист карточки называется её id (kg-skazka). Разделы идут в порядке карточки в приложении: Основное → Плашки → Тексты → Медиа → Дорога → Сезон → Важно → Ночлег и еда → На месте → Сервисы → Бюджет → Отзывы → Рядом → Проверка → Перевод EN.', ''),
        ('Поля: подпись в колонке B, значение — в колонке C. Жёлтая подпись — обязательное поле. Выпадающие списки — у страны, публикации, статуса, категории, региона, плашек.', ''),
        ('Списки (фото, подъезды, ночлег, тропы…): зелёная строка с названием, под ней подписи колонок и записи. Новая запись — в пустую строку под таблицей; если нужно больше, вставьте строку (Ctrl+«+») внутри таблицы.', ''),
        ('Колонку A (серые служебные метки) не трогайте — по ним программа читает файл.', ''),
        ('', ''),
        ('Как добавить новое место', 'h2'),
        ('1. Правой кнопкой по листу «_шаблон» → «Переместить или скопировать» → «Создать копию».', ''),
        ('2. Переименуйте копию в id нового места: kg-… или alt-…, латиница и дефисы (например kg-ala-kul). В поле id впишите то же самое.', ''),
        ('3. Заполните жёлтые поля, добавьте фото (файлы — в папку assets/img проекта), подъезды, ночлег. Публикация = черновик, пока не готово.', ''),
        ('4. Добавьте строку в «Оглавление» (id со ссылкой) и тикет «Новое место» в «Тикеты».', ''),
        ('5. Когда готово — публикация = опубликовано. Разработчик запускает: npm run places:import (ошибки покажет с листом и строкой).', ''),
        ('', ''),
        ('Правила заполнения', 'h2'),
        ('Месяцы — номера через «;» (5;6;7;8;9). Теги — id из листа «Теги» через «;». Да/нет — словами. Координаты — десятичные градусы (42.1565).', ''),
        ('Списки фраз (чего не хватает, противоречия) — каждая с новой строки в одной ячейке (Alt+Enter).', ''),
        ('Км, цены, время — текстом, как в источнике («≈8 км», «от 1000 сом», «2–3 ч»). Год цены — число.', ''),
        ('Ничего не выдумываем: нет данных — оставьте пусто и заведите тикет «Дозаполнить данные».', ''),
        ('Перевод EN: колонка «English» в конце карточки; после вычитки поставьте «Перевод EN проверен = да» в разделе «Проверка».', ''),
        ('', ''),
        ('Для разработчика (парсер)', 'h2'),
        ('Карточки — листы с именем ^(kg|alt)-[a-z0-9-]+$. Колонка A: ключ поля → значение в C; «#section» — заголовок; «#table:имя» → «#keys» (ключи колонок с C) → «#labels» → строки с пустой A до следующей непустой A.', ''),
        ('Схема и проверки — tools/places_xlsx.py, чтение/запись этого формата — tools/places_cards.py, описание — docs/09_DATA_EXCEL.md.', ''),
    ]
    ws.column_dimensions['A'].width = 150
    for r, (t, kind) in enumerate(lines, 1):
        c = ws.cell(r, 1, t)
        c.font = Font(name=F, size=16 if kind == 'h' else 12 if kind == 'h2' else 10, bold=bool(kind))
        c.alignment = Alignment(wrap_text=True, vertical='top')


def read_extras(path):
    """Тикеты, словарь тегов и теги мест из прошлой версии файла — чтобы повторная выгрузка их не стёрла."""
    res = {'tickets': None, 'tag_dict': None, 'tags': {}}
    if not os.path.exists(path):
        return res
    wb = load_workbook(path, read_only=True)
    if 'Тикеты' in wb.sheetnames:
        res['tickets'] = list(X.iter_rows(wb['Тикеты']))
    if 'Теги' in wb.sheetnames:
        res['tag_dict'] = list(X.iter_rows(wb['Теги']))
    for name in wb.sheetnames:
        if ID_RX.fullmatch(name):
            fields, _ = parse_card(wb[name])
            if fields.get('tags'):
                res['tags'][name] = [t for t in X.s(fields['tags']).split(';') if t.strip()]
    return res


# ---------------------------------------------------------------- чтение
def parse_card(ws):
    """Лист карточки → (поля {ключ: значение}, таблицы {имя: [строки с '_row']})."""
    fields, tables = {}, {}
    cur = keys = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        a = row[0] if row else None
        a = a.strip() if isinstance(a, str) else a
        if a:
            if a.startswith('#table:'):
                cur, keys = a[7:], None
                tables.setdefault(cur, [])
            elif a == '#keys' and cur:
                keys = [k for k in row[2:]]
            elif a in ('#labels', '#section', '#title'):
                if a != '#labels':
                    cur = None
            elif a.startswith('#'):
                cur = None
            else:
                cur = None
                fields[a] = row[2] if len(row) > 2 else None
            continue
        if cur and keys:
            vals = {k: v for k, v in zip(keys, row[2:]) if k}
            if any(v not in (None, '') for v in vals.values()):
                vals['pos'] = row[1] if row[1] not in (None, '') else len(tables[cur]) + 1
                vals['_row'] = i
                tables[cur].append(vals)
    return fields, tables


def build(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    tables = {'places': [], 'translations_en': [], 'tags': []}
    errors = []
    if 'Теги' in wb.sheetnames:
        tables['tags'] = list(X.iter_rows(wb['Теги']))
    for name in wb.sheetnames:
        if not ID_RX.fullmatch(name):
            if name.startswith(('kg', 'alt')) and not name.startswith('_'):
                errors.append(f'лист «{name}»: имя не похоже на id карточки (kg-… / alt-…, латиница и дефис) — лист пропущен')
            continue
        fields, tbls = parse_card(wb[name])
        fid = X.s(fields.get('id')).strip()
        if fid and fid != name:
            errors.append(f'лист «{name}»: поле id = «{fid}», а лист называется иначе — должны совпадать')
        fields['id'] = fid or name
        fields['_where'] = f'лист {name}'
        tables['places'].append(fields)
        for tname, rows in tbls.items():
            for r in rows:
                r['place_id'] = fields['id']
                r['_where'] = f'лист {name}, строка {r.pop("_row")}'
                tables.setdefault(tname, []).append(r)
    places, tr, reviewed, errs, warnings = X.build_tables(tables)
    return places, tr, reviewed, errors + errs, warnings


def do_import(path, write=True):
    places, tr, reviewed, errors, warnings = build(path)
    for w in warnings[:40]:
        print('  ⚠', w)
    if len(warnings) > 40:
        print(f'  … и ещё предупреждений: {len(warnings) - 40}')
    if errors:
        for e in errors:
            print('  ✗', e)
        print(f'✗ ошибок: {len(errors)} — places.json не изменён')
        return False
    pub = X.to_json(places)
    if not write:
        print(f'✓ проверка пройдена: карточек {len(places)} (к публикации {len(pub)}), предупреждений {len(warnings)}')
        return True
    data = json.load(open(X.PLACES, encoding='utf-8'))
    data['places'] = pub
    data['meta']['count'] = len(pub)
    for c in data['meta']['countries']:
        data['meta']['countries'][c]['count'] = sum(1 for p in pub if p['c'] == c)
    data['meta']['built'] = date.today().isoformat()
    json.dump(data, open(X.PLACES, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
    os.makedirs(X.TR_EN, exist_ok=True)
    for pid, m in tr.items():
        json.dump(m, open(os.path.join(X.TR_EN, f'{pid}.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    json.dump(sorted(reviewed), open(os.path.join(X.TR_EN, '_reviewed.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'✓ places.json: {len(pub)} мест · переводов EN: {len(tr)} · проверено редактором: {len(reviewed)}')
    return True


def roundtrip():
    import tempfile
    tmp = os.path.join(tempfile.gettempdir(), 'voyz-cards-roundtrip.xlsx')
    if os.path.exists(tmp):
        os.remove(tmp)
    export(tmp)
    places, tr, reviewed, errors, _ = build(tmp)
    assert not errors, errors[:5]
    orig = json.load(open(X.PLACES, encoding='utf-8'))['places']
    got = {p['id']: p for p in X.to_json(places)}
    diffs = 0
    for o in orig:
        g = dict(got[o['id']])
        g.pop('tags', None)
        if json.dumps(o, ensure_ascii=False, sort_keys=True) != json.dumps(g, ensure_ascii=False, sort_keys=True):
            diffs += 1
            for k in o:
                if json.dumps(o[k], sort_keys=True, ensure_ascii=False) != json.dumps(g.get(k), sort_keys=True, ensure_ascii=False):
                    print('  отличие', o['id'], k, str(o[k])[:100], '→', str(g.get(k))[:100])
                    break
    tr_bad = sum(1 for pid in got if os.path.exists(os.path.join(X.TR_EN, f'{pid}.json'))
                 and json.load(open(os.path.join(X.TR_EN, f'{pid}.json'), encoding='utf-8')) != tr.get(pid, {}))
    print(f'roundtrip: карточек {len(orig)}, отличий {diffs}; переводов с отличиями {tr_bad}')
    return diffs == 0 and tr_bad == 0


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'export'
    path = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else X.DEFAULT_XLSX
    ok = {'export': lambda: export(path) or True, 'import': lambda: do_import(path), 'check': lambda: do_import(path, write=False), 'roundtrip': roundtrip}[cmd]()
    sys.exit(0 if ok else 1)
